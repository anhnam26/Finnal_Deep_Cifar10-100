"""Build a Vietnamese, evidence-backed submission bundle without running training."""
from __future__ import annotations

import ast
import csv
import hashlib
import html
import io
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import textwrap
from datetime import datetime, timedelta, timezone
import zipfile
import xml.etree.ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))  # First-party project code, never third-party packages.
from src.python_policy import ensure_selected_python

ensure_selected_python()
os.environ.setdefault("MPLCONFIGDIR", str(ROOT / ".cache" / "matplotlib"))
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
import numpy as np
import torch
from src import engine
from src.models import build_model, profile_model

OUT = ROOT / "report"
DATASETS = ("cifar10", "cifar100")
NAMES = {"cifar10": "CIFAR-10", "cifar100": "CIFAR-100"}
EXPECTED_TESTS = {"cifar10": (97.18, 9718), "cifar100": (82.98, 8298)}
BLUE, GREEN, PURPLE, ORANGE = "#176B9B", "#27866A", "#7952A0", "#BD6A21"
STAMP = datetime.now(timezone.utc).astimezone(timezone(timedelta(hours=7)))
torch.set_num_threads(2)
plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10,
                     "axes.spines.top": False, "axes.spines.right": False})


def load_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


def save_json(path, value):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def sha(blob):
    return hashlib.sha256(blob).hexdigest()


def write_csv(path, rows):
    with Path(path).open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def display(value):
    if isinstance(value, float):
        return f"{value:.4f}".rstrip("0").rstrip(".")
    return str(value)


def clock_duration(seconds):
    seconds = round(seconds)
    return f"{seconds // 3600} giờ {(seconds % 3600) // 60} phút {seconds % 60} giây"


def final_note(runs):
    scores = [f'{NAMES[d]} {runs[d]["final"]["test_accuracy"]:.2f}%'
              for d in DATASETS if runs[d]["final"]["test_accuracy"] is not None]
    if not scores:
        return "Retrain 50k chưa có kết quả final test trong snapshot này."
    return "Final test retrain 50k đã có: " + "; ".join(scores) + "."


def collect():
    runs = {}
    for dataset in DATASETS:
        run = ROOT / "checkpoints" / f"default_{dataset}_compactmix_sgd_lr0.1_seed42"
        result = load_json(run / "result.json")
        assert result["status"] == "complete" and result["phase"] == "selection"
        assert not result["synthetic"] and not result["subset"]
        history, config = result["history"], result["config"]
        assert [row["epoch"] for row in history] == list(range(1, 601))
        assert result["metadata"]["source_sha256"] == engine.source_fingerprint()
        blob = (run / "best.pt").read_bytes()
        digest = sha(blob)
        checkpoint = torch.load(io.BytesIO(blob), map_location="cpu", weights_only=False)
        assert checkpoint["config"] == config
        assert checkpoint["epoch"] == result["best_epoch"]
        assert checkpoint["accuracy"] == result["best_val_accuracy"]
        candidates = [load_json(path) for path in (ROOT / "experiments/results").glob(f"{dataset}_best_test*.json")]
        matches = [item for item in candidates if item["checkpoint_sha256"] == digest]
        assert matches, f"Missing measured test evidence for {dataset}"
        test = sorted(matches, key=lambda item: item["completed_at_utc"])[-1]
        assert test["source_sha256"] == engine.source_fingerprint()
        assert test["epoch"] == checkpoint["epoch"] and test["weight_kind"] == checkpoint["kind"]
        assert test["training_config"] == config and test["split"] == "official_test"
        matrix = np.asarray(test["test_metrics"]["confusion_matrix"], dtype=int)
        classes = 10 if dataset == "cifar10" else 100
        assert matrix.shape == (classes, classes) and matrix.sum() == 10000
        assert np.trace(matrix) == test["correct"] and test["samples"] == 10000
        assert 100 * np.trace(matrix) / matrix.sum() == test["test_metrics"]["accuracy"]
        assert (test["test_metrics"]["accuracy"], test["correct"]) == EXPECTED_TESTS[dataset]
        split = load_json(run / "split.json")
        assert len(split["training"]) == 45000 and len(split["validation"]) == 5000
        assert set(split["training"]).isdisjoint(split["validation"])
        assert sorted(split["training"] + split["validation"]) == list(range(50000))
        labels = test.get("class_names", ["airplane", "automobile", "bird", "cat", "deer",
                                          "dog", "frog", "horse", "ship", "truck"])
        assert len(labels) == classes
        destination = OUT / "evidence" / dataset
        destination.mkdir(parents=True, exist_ok=True)
        for filename in ("config.json", "metadata.json", "metrics.csv", "training.log", "result.json", "split.json"):
            shutil.copyfile(run / filename, destination / filename)
        save_json(destination / "test_evaluation.json", test)
        weight_dir = OUT / "checkpoints" / dataset
        weight_dir.mkdir(parents=True, exist_ok=True)
        (weight_dir / "best.pt").write_bytes(blob)
        save_json(weight_dir / "checkpoint_info.json", {
            "checkpoint_sha256": digest, "epoch": checkpoint["epoch"], "kind": checkpoint["kind"],
            "validation_accuracy": checkpoint["accuracy"], "train_images": 45000, "config": config})
        model = build_model(config["model"], classes).eval()
        model.load_state_dict(checkpoint["model"], strict=True)
        profile = profile_model(model)
        assert profile == result["profile"]
        layers = trace_layers(model)
        assert sum(row["parameters"] for row in layers) == profile["parameters"]
        write_csv(destination / "layers.csv", layers)
        (destination / "model_repr.txt").write_text(str(model), encoding="utf-8")
        save_json(destination / "verified_profile.json", profile)
        per_class = class_metrics(matrix, labels)
        assert np.all(matrix.sum(1) == 10000 // classes)
        assert abs(np.mean([row["f1_percent"]/100 for row in per_class])
                   - test["test_metrics"]["macro_f1"]) < 1e-6
        write_csv(destination / "per_class_metrics.csv", per_class)
        with (destination / "confusion_matrix.csv").open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(["true/predicted"] + labels)
            for label, row in zip(labels, matrix):
                writer.writerow([label] + row.tolist())
        best_row = history[checkpoint["epoch"] - 1]
        assert best_row["ema_val_accuracy" if checkpoint["kind"] == "ema" else "val_accuracy"] == checkpoint["accuracy"]
        final_dir = run.with_name(run.name + "_final")
        final = {"status": "Chưa bắt đầu", "completed_epochs": 0,
                 "planned_epochs": result["best_epoch"], "test_accuracy": None}
        if (final_dir / "frozen.json").exists():
            final["status"] = "Chưa có kết quả hoàn tất"
            final["freeze"] = load_json(final_dir / "frozen.json")
            if (final_dir / "metrics.csv").exists():
                # Atomic CSV replaces permit a consistent snapshot without locking the trainer.
                blob_csv = (final_dir / "metrics.csv").read_bytes()
                rows = list(csv.DictReader(io.StringIO(blob_csv.decode("utf-8-sig"))))
                final["completed_epochs"] = int(rows[-1]["epoch"]) if rows else 0
                (destination / "final_progress.csv").write_bytes(blob_csv)
            if (final_dir / "result.json").exists():
                final_result = load_json(final_dir / "result.json")
                final.update(status=final_result["status"], test_accuracy=final_result["test_accuracy"])
        save_json(destination / "final_status_snapshot.json", final)
        runs[dataset] = {"result": result, "test": test, "config": config, "history": history,
                         "profile": profile, "layers": layers, "per_class": per_class,
                         "best_row": best_row, "labels": labels, "matrix": matrix,
                         "sha256": digest, "final": final}
    return runs


def trace_layers(model):
    rows, handles = [], []
    def hook(name):
        def record(module, inputs, output):
            kernel = list(module.kernel_size) if isinstance(module, torch.nn.Conv2d) else ""
            groups = module.groups if isinstance(module, torch.nn.Conv2d) else ""
            kind = type(module).__name__
            if isinstance(module, torch.nn.Conv2d):
                kind = "depthwise" if groups == module.in_channels == module.out_channels else "grouped" if groups > 1 else "dense/pointwise"
            rows.append({"name": name, "type": type(module).__name__, "conv_kind": kind,
                         "input_shape": str(list(inputs[0].shape)),
                         "output_shape": str(list(output.shape)),
                         "kernel": str(kernel), "stride": str(list(module.stride)) if kernel else "",
                         "padding": str(list(module.padding)) if kernel else "", "groups": groups,
                         "parameters": sum(p.numel() for p in module.parameters(recurse=False))})
        return record
    for name, module in model.named_modules():
        if not list(module.children()):
            handles.append(module.register_forward_hook(hook(name)))
    try:
        with torch.no_grad():
            assert model(torch.zeros(1, 3, 32, 32)).shape[-1] in (10, 100)
    finally:
        for handle in handles:
            handle.remove()
    return rows


def class_metrics(matrix, labels):
    tp = np.diag(matrix)
    precision = np.divide(tp, matrix.sum(0), out=np.zeros(len(tp), dtype=float), where=matrix.sum(0) > 0)
    recall = tp / matrix.sum(1)
    f1 = np.divide(2 * precision * recall, precision + recall,
                   out=np.zeros(len(tp), dtype=float), where=(precision + recall) > 0)
    return [{"class_id": i, "class_name": label, "support": int(matrix[i].sum()),
             "correct": int(tp[i]), "precision_percent": float(precision[i] * 100),
             "recall_percent": float(recall[i] * 100), "f1_percent": float(f1[i] * 100)}
            for i, label in enumerate(labels)]


def box(ax, x, y, width, height, label, color=BLUE, size=11):
    patch = FancyBboxPatch((x-width/2, y-height/2), width, height,
                           boxstyle="round,pad=0.03,rounding_size=0.08",
                           linewidth=1.3, edgecolor=color, facecolor=color + "12")
    ax.add_patch(patch)
    ax.text(x, y, label, ha="center", va="center", fontsize=size, color="#182B3A")


def arrow(ax, start, end, color="#536575"):
    ax.annotate("", xy=end, xytext=start,
                arrowprops={"arrowstyle": "-|>", "lw": 1.4, "color": color})


def diagram(name, figsize, limits):
    fig, ax = plt.subplots(figsize=figsize)
    ax.set_xlim(*limits[0]); ax.set_ylim(*limits[1]); ax.axis("off")
    return fig, ax


def save_figure(fig, name):
    directory = OUT / "figures"
    directory.mkdir(exist_ok=True)
    fig.savefig(directory / f"{name}.png", dpi=190, bbox_inches="tight", facecolor="white")
    fig.savefig(directory / f"{name}.svg", bbox_inches="tight", facecolor="white")
    plt.close(fig)


def architecture_diagrams():
    fig, ax = diagram("architecture", (8.4, 11), ((0, 10), (0, 12)))
    nodes = [
        (11.25, .65, "INPUT RGB\nN × 3 × 32 × 32", BLUE),
        (10.1, .75, "STEM · Conv 3×3, stride 1, padding 1\n3 → 40 · N × 40 × 32 × 32", BLUE),
        (8.9, .85, "STAGE 1 · 2 MixBlock\nC=40 · H=W=32 · expansion=2", GREEN),
        (7.6, .85, "TRANSITION 1 · BN → SiLU → AvgPool 2×2\nConv 1×1: 40 → 112 · N × 112 × 16 × 16", ORANGE),
        (6.3, .85, "STAGE 2 · 3 MixBlock\nC=112 · H=W=16 · expansion=2", GREEN),
        (5.0, .85, "TRANSITION 2 · BN → SiLU → AvgPool 2×2\nConv 1×1: 112 → 256 · N × 256 × 8 × 8", ORANGE),
        (3.7, .85, "STAGE 3 · 8 MixBlock\nC=256 · H=W=8 · expansion=2", GREEN),
        (2.15, 1.1, "HEAD · BN(256) → SiLU → Global AvgPool\nFlatten → Dropout(p=0) → Linear(256, K)\nK=10 / K=100 · logits N × K", PURPLE),
        (.65, .65, "DỰ ĐOÁN · argmax(logits)\nSoftmax chỉ cần khi hiển thị xác suất", PURPLE)]
    for y, height, label, color in nodes:
        box(ax, 5, y, 8.6, height, label, color)
    for first, second in zip(nodes, nodes[1:]):
        arrow(ax, (5, first[0]-first[1]/2), (5, second[0]+second[1]/2))
    save_figure(fig, "01_compactmix_overview")

    fig, ax = diagram("mixblock", (9.5, 13), ((0, 11), (-.3, 14.2)))
    chain = [
        (13.45, .65, "x: N × C × H × W", BLUE),
        (12.1, 1.05, "PRE · BN(C) → SiLU → Conv 1×1: C → 2C\nBN(2C) → SiLU · chia đều 2 nhánh", BLUE)]
    for y, height, label, color in chain:
        box(ax, 6, y, 8.6, height, label, color)
    arrow(ax, (6, 13.12), (6, 12.65))
    box(ax, 3.55, 10.4, 3.7, 1.1, "LOCAL · C kênh\nDW Conv 3×3\npadding=1, groups=C", GREEN)
    box(ax, 8.4, 10.4, 3.7, 1.1, "CONTEXT · C kênh\nDW Conv 5×5\npadding=2, groups=C", PURPLE)
    arrow(ax, (4.2, 11.55), (3.55, 10.98))
    arrow(ax, (7.8, 11.55), (8.4, 10.98))
    box(ax, 6, 8.9, 8.6, .75, "Concat theo chiều kênh → N × 2C × H × W", BLUE)
    arrow(ax, (3.55, 9.82), (4.2, 9.32)); arrow(ax, (8.4, 9.82), (7.8, 9.32))
    tails = [
        (7.65, .9, "BN(2C) → SiLU → Grouped Conv 3×3\n2C → 2C · groups=8 · padding=1", ORANGE),
        (6.3, .75, "BN(2C) → SiLU", BLUE),
        (5.1, .8, "SE ChannelGate · attention theo kênh\nGAP → Conv 1×1 → SiLU → Conv 1×1 → Sigmoid", PURPLE),
        (3.8, .75, "Conv 1×1 · chiếu 2C → C", BLUE),
        (2.55, .75, "DropPath · p tăng từ 0 đến 0,1 theo 13 block\nKhông hoạt động khi eval()", ORANGE),
        (1.1, .8, "CỘNG RESIDUAL · y = x + DropPath(residual)\nN × C × H × W", GREEN)]
    previous = (8.9, .75)
    for y, height, label, color in tails:
        box(ax, 6, y, 8.6, height, label, color, 10.5)
        arrow(ax, (6, previous[0]-previous[1]/2), (6, y+height/2))
        previous = (y, height)
    ax.plot([1.67, .55, .55, 1.35], [13.45, 13.45, 1.1, 1.1], color=GREEN, lw=1.6)
    arrow(ax, (1.35, 1.1), (1.67, 1.1), GREEN)
    ax.text(.25, 7.2, "identity / skip connection", rotation=90,
            va="center", ha="center", color=GREEN, fontsize=11)
    save_figure(fig, "02_mixblock")

    fig, ax = diagram("se", (11, 3.5), ((0, 13.2), (0, 3)))
    parts = [(1.15, "Feature\n2C × H × W"), (3.65, "Global\nAvgPool\n2C × 1 × 1"),
             (6.2, "Conv 1×1\n2C → d\nSiLU"), (8.75, "Conv 1×1\nd → 2C\nSigmoid"),
             (11.35, "Nhân theo kênh\nx × gate\n2C × H × W")]
    for x, label in parts:
        box(ax, x, 1.8, 2.1, 1.25, label, PURPLE, 10.5)
    for first, second in zip(parts, parts[1:]):
        arrow(ax, (first[0]+1.1, 1.8), (second[0]-1.1, 1.8))
    ax.plot([1.15, 1.15, 11.35], [2.46, 2.8, 2.8], color=GREEN, lw=1.4)
    arrow(ax, (11.35, 2.8), (11.35, 2.45), GREEN)
    ax.text(6.5, .5, "d = max(8, floor(2C / 16)) · d = 8 / 14 / 32 tại Stage 1 / 2 / 3",
            ha="center", fontsize=11)
    save_figure(fig, "03_se_gate")

    fig, ax = diagram("protocol", (11, 6.7), ((0, 12), (0, 7)))
    box(ax, 2, 6.3, 3.6, .8, "Official training: 50.000 ảnh", BLUE)
    box(ax, 2, 4.7, 3.6, 1, "Stratified split, seed 2026\n45.000 train / 5.000 validation", BLUE)
    arrow(ax, (2, 5.87), (2, 5.25))
    box(ax, 7.8, 4.7, 6.3, 1, "600 epoch · train raw + EMA\nValidation chọn best.pt · không chọn theo test", GREEN)
    arrow(ax, (3.85, 4.7), (4.58, 4.7))
    box(ax, 7.8, 2.85, 6.3, 1, "TEST ĐÃ ĐO · checkpoint best của train 45k\n10.000 ảnh test · theo yêu cầu người dùng", PURPLE)
    arrow(ax, (7.8, 4.15), (7.8, 3.4))
    box(ax, 2, 2.85, 3.6, 1, "Official test\n10.000 ảnh / dataset", PURPLE)
    arrow(ax, (3.85, 2.85), (4.58, 2.85))
    box(ax, 6, .75, 11.5, 1, "GIAI ĐOẠN RIÊNG · retrain mới trên 50k, cố định recipe + số epoch\n"
        "CIFAR-10: 567 epoch / CIFAR-100: 399 epoch · test final sau khi hoàn tất", ORANGE)
    save_figure(fig, "04_data_evaluation_protocol")


def result_figures(runs):
    for dataset, item in runs.items():
        history, best_epoch = item["history"], item["result"]["best_epoch"]
        epochs = [r["epoch"] for r in history]
        fig, axes = plt.subplots(2, 2, figsize=(12, 8), constrained_layout=True)
        for key, label, color in [("val_accuracy", "Validation raw", BLUE),
                                  ("ema_val_accuracy", "Validation EMA", GREEN),
                                  ("best_val_accuracy", "Best tích lũy", PURPLE)]:
            axes[0, 0].plot(epochs, [r[key] for r in history], label=label, color=color, lw=1.2)
        axes[0, 0].set(title="Validation top-1", ylabel="Accuracy (%)")
        axes[0, 0].legend(fontsize=8)
        axes[0, 1].plot(epochs, [r["train_mixed_accuracy"] for r in history], color=ORANGE, lw=1.1)
        axes[0, 1].set(title="Train mixed-label accuracy", ylabel="Accuracy có trọng số (%)")
        for key, label, color in [("train_loss", "Train soft-target loss", ORANGE),
                                  ("val_loss", "Validation raw CE", BLUE),
                                  ("ema_val_loss", "Validation EMA CE", GREEN)]:
            axes[1, 0].plot(epochs, [r[key] for r in history], label=label, color=color, lw=1.1)
        axes[1, 0].set(title="Loss", ylabel="Loss")
        axes[1, 0].legend(fontsize=8)
        axes[1, 1].plot(epochs, [r["lr"] for r in history], color=BLUE)
        axes[1, 1].set(title="Learning rate ở batch cuối epoch", ylabel="LR (log)", yscale="log")
        for axis in axes.flat:
            axis.axvline(best_epoch, ls="--", color=PURPLE, alpha=.6, lw=1)
            axis.set_xlabel("Epoch")
            axis.grid(alpha=.15)
        fig.suptitle(f"{NAMES[dataset]} · 600 epoch · best tại epoch {best_epoch}", fontsize=15)
        save_figure(fig, f"{dataset}_training_curves")
        matrix = item["matrix"]
        normalized = matrix / matrix.sum(1, keepdims=True) * 100
        fig, ax = plt.subplots(figsize=(9, 8), constrained_layout=True)
        im = ax.imshow(normalized, cmap="Blues", vmin=0, vmax=100)
        ax.set(title=f"{NAMES[dataset]} · Confusion matrix trên 10.000 ảnh test",
               xlabel="Nhãn dự đoán", ylabel="Nhãn thật")
        if dataset == "cifar10":
            ax.set_xticks(range(10), item["labels"], rotation=35, ha="right", fontsize=9)
            ax.set_yticks(range(10), item["labels"], fontsize=9)
            for row in range(10):
                for column in range(10):
                    ax.text(column, row, str(matrix[row, column]), ha="center", va="center",
                            fontsize=8, color="white" if normalized[row, column] > 50 else "#19364A")
        else:
            ax.set_xticks(range(0, 100, 10)); ax.set_yticks(range(0, 100, 10))
        fig.colorbar(im, ax=ax, label="Tỷ lệ theo nhãn thật (%)", shrink=.8)
        save_figure(fig, f"{dataset}_confusion_matrix")
        metrics = sorted(item["per_class"], key=lambda row: row["recall_percent"])[:20]
        fig, ax = plt.subplots(figsize=(10, 7), constrained_layout=True)
        ax.barh([row["class_name"] for row in metrics], [row["recall_percent"] for row in metrics], color=BLUE)
        ax.invert_yaxis(); ax.set_xlim(0, 105)
        ax.set(title=f"{NAMES[dataset]} · Recall theo lớp" + (" · 20 lớp thấp nhất" if dataset == "cifar100" else ""),
               xlabel="Recall (%) trên tập test")
        for i, row in enumerate(metrics):
            ax.text(row["recall_percent"] + .6, i, f'{row["recall_percent"]:.1f}%', va="center", fontsize=9)
        save_figure(fig, f"{dataset}_class_recall")


def source_segment(filename, symbol):
    source = (ROOT / filename).read_text(encoding="utf-8")
    for node in ast.parse(source).body:
        if isinstance(node, (ast.ClassDef, ast.FunctionDef, ast.AsyncFunctionDef)) and node.name == symbol:
            first_line = min([node.lineno] + [item.lineno for item in node.decorator_list])
            return "\n".join(source.splitlines()[first_line-1:node.end_lineno])
    raise ValueError(f"Missing source symbol {filename}:{symbol}")


def make_content(runs):
    content = []
    def heading(title, page=False): content.append(("heading", title, page))
    def para(text): content.append(("paragraph", text))
    def table(headers, rows): content.append(("table", headers, rows))
    def image(name, caption): content.append(("image", f"figures/{name}.png", caption))
    def code(filename, symbol):
        content.append(("code", f"{filename} — {symbol}", source_segment(filename, symbol)))
    heading("Tóm tắt kết quả đã hoàn thành")
    para("Đề tài: thiết kế và huấn luyện mô hình nhẹ CompactMix-L từ đầu trên CIFAR-10 và CIFAR-100. "
         "Hai thí nghiệm selection đã hoàn thành 600 epoch trên 45.000 ảnh train và 5.000 ảnh validation. "
         "Các checkpoint best được chọn theo validation; tập test chính thức có 10.000 ảnh mỗi dataset.")
    table(["Thông tin", "CIFAR-10", "CIFAR-100"], [
        ["Best validation top-1", "97,56%", "83,36%"],
        ["Test top-1 đã đo (best train 45k)", "97,18% (9.718/10.000)", "82,98% (8.298/10.000)"],
        ["Epoch / loại trọng số", "567 / EMA", "399 / EMA"],
        ["Tham số học được", "5.248.172", "5.271.302"],
        ["FLOPs ước lượng / ảnh", "842.406.944", "842.453.114"],
        ["Optimizer / LR ban đầu", "SGD / 0,1", "SGD / 0,1"]])
    para("Các giá trị test trên thuộc checkpoint được train trên 45k, được đo theo yêu cầu người dùng. "
         "Chúng không phải kết quả của bản retrain trên 50k. "
         "Không có pretrained, ensemble, test-time augmentation hay dữ liệu ngoài trong hai lần đo.")
    para(f"Báo cáo là ảnh chụp trạng thái tại {STAMP:%d/%m/%Y %H:%M:%S} (UTC+7). "
         "Mã nguồn, log, cấu hình, chỉ số theo lớp, checkpoint best và hash đi kèm trong thư mục report.")
    heading("1. Đề bài và mức độ đáp ứng", True)
    para("Nguồn yêu cầu: Final exam.docx, đính kèm nguyên bản trong evidence/. "
         "Đề yêu cầu một model L nhẹ, không quá 6M tham số, FLOPs dưới 1G; "
         "train/test trên cả hai dataset với nhiều learning rate ban đầu và cả SGD, Adam; nộp báo cáo PDF.")
    table(["Yêu cầu", "Bằng chứng / trạng thái"], [
        ["<= 6.000.000 tham số", "Đạt cho cả hai head; đếm trực tiếp model.parameters()."],
        ["< 1.000.000.000 FLOPs", "Đạt theo phép đếm inference được mô tả tại mục 4."],
        ["Train và test CIFAR-10 / CIFAR-100", "Đã hoàn thành selection 45k/5k và test best trên đủ 10k ảnh mỗi dataset."],
        ["So sánh SGD, Adam và nhiều LR", "CHƯA HOÀN THÀNH: mới có hai run full-data SGD, LR=0,1."],
        ["Báo cáo PDF", "Bao_cao_CIFAR10_CIFAR100.pdf trong gói này."],
        ["Tính khác biệt kiến trúc giữa nhóm", "Tự triển khai kết hợp các khối chuẩn; chưa có bằng chứng kiểm tra trùng với nhóm khác."]])
    para("Không điền các điểm số còn thiếu bằng mục tiêu. Các mục tiêu nghiên cứu 98% / 90% không phải "
         "điểm số đã đạt, cũng không phải ngưỡng bắt buộc trong văn bản đề bài.")
    para("Tiêu chí chấm: vấn đáp 30%, hiệu năng CIFAR-10 30%, hiệu năng CIFAR-100 30%, chất lượng báo cáo 10%. "
         "Đề lưu ý phải vượt qua vấn đáp; kiến trúc không được trùng nhóm khác.")
    heading("2. Dữ liệu và giao thức thực nghiệm", True)
    table(["Thuộc tính", "CIFAR-10", "CIFAR-100"], [
        ["Ảnh đầu vào / nhãn", "RGB 32×32 / 10 lớp", "RGB 32×32 / 100 fine classes"],
        ["Official train / test", "50.000 / 10.000", "50.000 / 10.000"],
        ["Selection train / validation", "45.000 / 5.000", "45.000 / 5.000"],
        ["Số train / val / test mỗi lớp", "4.500 / 500 / 1.000", "450 / 50 / 100"],
        ["Seed khởi tạo / split", "42 / 2026", "42 / 2026"],
        ["Mean RGB", "0,4914; 0,4822; 0,4465", "0,5071; 0,4867; 0,4408"],
        ["Std RGB", "0,2470; 0,2435; 0,2616", "0,2675; 0,2565; 0,2761"]])
    image("04_data_evaluation_protocol", "Sơ đồ dữ liệu: tách selection và retrain final; test 45k được đo riêng.")
    para("Split được lấy phân tầng từ official training, không giao nhau và phủ đủ 50.000 chỉ số. "
         "File split.json lưu đầy đủ index cùng SHA256, có thể kiểm tra lại từng ảnh thuộc phần nào.")
    para("Train sử dụng RandomCrop(32, padding=4, padding_mode='reflect'), RandomHorizontalFlip, "
         "AutoAugment policy CIFAR10 trên cả hai dataset, rồi ToTensor và Normalize. "
         "Validation/test chỉ ToTensor và Normalize; không crop/flip ngẫu nhiên.")
    para("Chỉ số đánh giá: top-1=100·số dự đoán đúng/N; top-5=100·số mẫu có nhãn thật trong 5 logits cao nhất/N. "
         "Theo từng lớp: precision=TP/(TP+FP), recall=TP/(TP+FN), F1=2·precision·recall/(precision+recall). "
         "Macro-F1 là trung bình F1 không trọng số trên K lớp. Confusion matrix dùng hàng là nhãn thật, "
         "cột là nhãn dự đoán; mỗi ảnh được đánh giá đúng một lần.")
    para("Ban đầu chọn kiến trúc, recipe và best bằng validation. Sau selection, hai bản best 45k đã được "
         "đánh giá test theo yêu cầu người dùng. Vì test đã được mở, mọi thay đổi tiếp theo dựa trên điểm "
         "test sẽ không còn đáp ứng giao thức test giữ kín ban đầu. Báo cáo này không tuyên bố đã thực hiện "
         "ablation hay tối ưu thêm dựa trên test.")
    heading("3. Kiến trúc tổng thể CompactMix-L", True)
    image("01_compactmix_overview", "Sơ đồ model thực tế; cùng backbone, head khác số lớp.")
    table(["Khối", "Output (bỏ chiều batch)", "Chi tiết"], [
        ["Input", "3×32×32", "RGB"],
        ["Stem", "40×32×32", "Conv 3×3, stride=1, padding=1, không bias"],
        ["Stage 1", "40×32×32", "2 MixBlock"],
        ["Transition 1", "112×16×16", "BN40 → SiLU → AvgPool2 → Conv1×1(40,112)"],
        ["Stage 2", "112×16×16", "3 MixBlock"],
        ["Transition 2", "256×8×8", "BN112 → SiLU → AvgPool2 → Conv1×1(112,256)"],
        ["Stage 3", "256×8×8", "8 MixBlock"],
        ["Head", "256 → K", "BN256 → SiLU → GAP → Dropout(0) → Linear"],
        ["Logits", "K=10 / K=100", "Không softmax trong forward; argmax cho top-1"]])
    para("Stem giữ nguyên độ phân giải 32×32; model chỉ giảm kích thước không gian hai lần. "
         "Khối cuối rộng hơn để biểu diễn thông tin khi H,W nhỏ. Ý tưởng là phối hợp nhiều thang "
         "không gian trong một block và chia sẻ cùng backbone cho hai dataset.")
    heading("3.1. Cấu trúc MixBlock", True)
    image("02_mixblock", "Luồng tensor trong MixBlock; C giữ nguyên trước và sau residual.")
    table(["Stage", "C / 2C / H=W", "Grouped conv", "SE hidden d"], [
        ["1", "40 / 80 / 32", "groups=8; 10 kênh/group", "8"],
        ["2", "112 / 224 / 16", "groups=8; 28 kênh/group", "14"],
        ["3", "256 / 512 / 8", "groups=8; 64 kênh/group", "32"]])
    para("Sau Conv1×1 mở rộng C→2C, chunk chia thành hai tensor C kênh. Nhánh local dùng depthwise 3×3; "
         "nhánh context dùng depthwise 5×5. Concat trả về 2C kênh; grouped 3×3 trộn theo 8 nhóm. "
         "SE tạo trọng số kênh; Conv1×1 cuối chiếu về C để cộng với x.")
    para("Công thức forward: u=Pre(x); (a,b)=chunk(u,2); "
         "r=Project(SE(ActNorm(GroupConv(Concat(DW3(a),DW5(b)))))); "
         "y=x+DropPath(r). Phương trình r là mô tả rút gọn; thứ tự BN/SiLU đầy đủ nằm trong sơ đồ và code.")
    para("Depthwise giảm chi phí xử lý không gian nhưng không tự trộn kênh; grouped convolution trộn trong "
         "từng nhóm, còn Conv1×1 dense trộn tất cả kênh. Vì không có ablation full-data, chưa thể định lượng "
         "riêng mức cải thiện accuracy do nhánh 5×5, SE hoặc DropPath.")
    heading("3.2. SE, residual và các kiểu convolution", True)
    image("03_se_gate", "SE xử lý tensor 2C kênh, không phải C kênh.")
    para("SE: z=GAP(x), g=sigmoid(W2·SiLU(W1·z)), output=x·g; "
         "d=max(8,floor(2C/16)). Hai Conv1×1 của SE có bias. Sigmoid cho hệ số trong (0,1).")
    table(["Kiểu Conv", "Cách hoạt động / vai trò"], [
        ["Dense 3×3", "Stem: mỗi output channel kết nối với tất cả input channels."],
        ["Pointwise 1×1", "Mở rộng/chiếu kênh, transition và hai lớp SE."],
        ["Depthwise 3×3 / 5×5", "groups=C trên tensor nhánh C kênh; một filter trên mỗi kênh."],
        ["Grouped 3×3", "8 nhóm trên tensor 2C kênh; giữ H,W bằng padding=1."],
        ["Residual identity", "Đường x không đổi kích thước; cộng với nhánh r."],
        ["DropPath", "Trong train: bỏ toàn bộ residual của một mẫu theo xác suất p; eval: identity."]])
    para("13 block có drop probability p_i=0,1·i/12 (i=0..12). residual_scale=1, "
         "context_kernel=5, expansion=2, attention='se', norm='bn', activation='silu', "
         "stem='simple', pool='avg', dropout=0 là các mặc định đã dùng. "
         "BN sử dụng eps=1e-5 và momentum=0,1 theo mặc định PyTorch.")
    para("Model hỗ trợ tùy chọn ECA/không attention, BN/GN/GroupNorm(1,C), kernel context 3/5/7, "
         "pool avgmax và stacked stem; các biến thể này chưa có kết quả full-data trong báo cáo. "
         "Các mạng ResidualNet, PreActBlock, CIFARConvNeXt và adapter timm có source đi kèm, "
         "nhưng không phải model tạo ra các điểm số đang báo cáo.")
    heading("4. Tham số, FLOPs và bảng layer", True)
    table(["Dataset", "Parameters", "MACs", "FLOPs", "GFLOPs"], [
        [NAMES[d], f'{runs[d]["profile"]["parameters"]:,}', f'{runs[d]["profile"]["macs"]:,}',
         f'{runs[d]["profile"]["flops"]:,}', f'{runs[d]["profile"]["flops"]/1e9:.6f}'] for d in DATASETS])
    para("Đã tạo lại model trên CPU, load đúng best weights, forward input [1,3,32,32], "
         "đếm lại parameters và profile; các giá trị trùng metadata gốc. "
         "Bảng đầy đủ từng layer, input/output shape, kernel, stride, padding, groups và parameters "
         "nằm ở evidence/cifar10/layers.csv và evidence/cifar100/layers.csv.")
    para("Conv2d: P=Cout·(Cin/groups)·kh·kw + bias (nếu có). "
         "MAC=N·Hout·Wout·Cout·(Cin/groups)·kh·kw. Linear: P=Cin·Cout+Cout, MAC=N·Cin·Cout. "
         "BN có gamma và beta, nên 2C tham số học được; running mean/variance là buffer.")
    para("Ví dụ một block Stage 1 (C=40): pre Conv1×1 có 40·80=3.200 tham số; "
         "local DW3 có 40·9=360; context DW5 có 40·25=1.000; grouped3 có 80·(80/8)·9=7.200; "
         "projection1 có 80·40=3.200. SE: (80·8+8)+(8·80+80)=1.368; "
         "bốn BN đóng góp 2·40+3·2·80=560. Tổng block: 16.888 tham số.")
    para("Head Linear: CIFAR-10 có 256·10+10=2.570 tham số; CIFAR-100 có 25.700. "
         "Chênh lệch toàn model là 23.130 tham số; backbone giống nhau.")
    para("Quy ước: một MAC=2 FLOPs; cộng bias, BN (4/element), SiLU (6), Sigmoid (4), "
         "pooling, channel multiply và residual addition. Các thao tác split/cat/reshape "
         "không tính là FLOPs; DropPath/Dropout là identity khi eval. Đây là ước lượng phép toán "
         "inference, không phải latency phần cứng, training FLOPs hay phép đo năng lượng.")
    heading("5. Cấu hình train và hàm loss", True)
    config = runs["cifar10"]["config"]
    table(["Thiết lập", "Giá trị thực tế ở cả hai selection run"], [
        ["Epoch / batch / workers", "600 / 128 / 2"],
        ["Optimizer", "SGD; momentum=0,9; Nesterov=True"],
        ["LR / schedule", "0,1; warmup 5 epoch; cosine; min factor 0,001 (=LR 0,0001)"],
        ["Weight decay", "0,0005 cho tensor ndim>1; 0 cho bias/norm ndim<=1"],
        ["Loss / label smoothing", "Soft-target cross entropy; smoothing=0,05"],
        ["Mixup / CutMix", "alpha=0,2 / 1,0; mixing probability=0,75"],
        ["Mixup vs CutMix", "Khi mixing được chọn, chọn CutMix với xác suất 0,5"],
        ["Mixing cuối quá trình", "Tắt ở 30 epoch cuối (571..600) của selection"],
        ["EMA", "Decay=min(0,999, (1+global_step)/(10+global_step))"],
        ["Gradient clipping / AMP", "Global grad norm<=5; CUDA float16 autocast, GradScaler init=4096"],
        ["Determinism", "torch deterministic=True; TF32=False; seed per image/epoch"],
        ["Random erasing / Dropout head", "0 / 0"],
        ["Khởi tạo", "Conv Kaiming normal; Linear trunc_normal(std=0,02); bias=0"]])
    para("LR được cập nhật mỗi batch với progress=e+b/B, e đánh số 0..599 trong engine. "
         "Trong warmup: lr=0,1·(0,1+0,9·progress/5). Sau warmup: "
         "lr=0,1·[0,001+0,999·(1+cos(pi·fraction))/2], fraction=(progress-5)/(600-5). "
         "Log lưu LR ở batch cuối, nên epoch 1 có LR khoảng 0,02795, không phải LR đầu tiên của batch 1.")
    para("Targets: y_mix=lambda·onehot(y)+(1-lambda)·onehot(y_perm); "
         "q=(1-eps)·y_mix+eps/K; loss=-mean(sum_k q_k·log_softmax(logits)_k). "
         "CutMix hiệu chỉnh lambda theo diện tích patch thực tế sau khi cắt ở biên ảnh.")
    para("Train mixed accuracy=100·[lambda·đúng với y+(1-lambda)·đúng với y_perm]/N. "
         "Đây là chỉ số có trọng số của nhãn trộn, không phải accuracy train sạch. "
         "Train loss dùng soft labels/smoothing, validation/test loss dùng cross entropy nhãn cứng; "
         "không so sánh hai loss như cùng một bài toán hoàn toàn.")
    para("EMA chỉ cập nhật sau optimizer step thành công; buffer nguyên được copy. "
         "Nếu GradScaler skip update, global_step và EMA không tiến. "
         "Vì raw và EMA đều được kiểm tra validation, best có thể thuộc một trong hai loại; "
         "hai best hiện tại đều là EMA.")
    heading("6. Kết quả CIFAR-10", True)
    dataset_section(content, runs["cifar10"], "cifar10")
    heading("7. Kết quả CIFAR-100", True)
    dataset_section(content, runs["cifar100"], "cifar100")
    heading("8. So sánh hai dataset và phần còn thiếu", True)
    table(["Chỉ số", "CIFAR-10", "CIFAR-100"], [
        ["Best val → test", "97,56 → 97,18%", "83,36 → 82,98%"],
        ["Val–test gap", "0,38 điểm phần trăm", "0,38 điểm phần trăm"],
        ["Số dự đoán sai / 10k test", "282", "1.702"],
        ["Best epoch / total", "567 / 600", "399 / 600"],
        ["Số lớp / head", "10 / Linear(256,10)", "100 / Linear(256,100)"]])
    para("CIFAR-100 có nhiều lớp hơn và ít ảnh trên mỗi lớp; kết quả thấp hơn dù backbone và recipe giống nhau. "
         "Điều này mô tả hai phép đo hiện tại, không xác lập nguyên nhân hay độ ưu việt của một component. "
         "Cần ablation và lặp seed để xác nhận các giả thuyết.")
    comparison = []
    for optimizer, lrs in (("SGD", (.05, .1, .15)), ("Adam", (.0003, .001, .003))):
        for lr in lrs:
            comparison.append([optimizer, str(lr),
                "97,56 / 97,18" if optimizer == "SGD" and lr == .1 else "Chưa chạy",
                "83,36 / 82,98" if optimizer == "SGD" and lr == .1 else "Chưa chạy"])
    table(["Optimizer", "Initial LR", "C10 val / test (%)", "C100 val / test (%)"], comparison)
    para("Các hàng chưa chạy là cấu hình hỗ trợ bởi --matrix, không phải các kết quả thực nghiệm. "
         "Chưa có full-data Adam, nhiều LR, multi-seed mean/std hoặc ablation kernel/SE/EMA/augmentation. "
         "Phần so sánh optimizer/LR bắt buộc trong đề còn thiếu; không tuyên bố hoàn thành toàn bộ yêu cầu.")
    heading("9. Retrain final trên 50.000 ảnh", True)
    table(["Dataset", "Số epoch đã freeze", "Epoch trong snapshot", "Kết quả final test"], [
        [NAMES[d], runs[d]["final"]["planned_epochs"], runs[d]["final"]["completed_epochs"],
         "Chưa có" if runs[d]["final"]["test_accuracy"] is None else f'{runs[d]["final"]["test_accuracy"]:.2f}%']
        for d in DATASETS])
    para("Recipe, loại trọng số và số epoch final được freeze từ validation winner. "
         "Final reinitialize model và train trên cả 50.000 ảnh, không fine-tune từ best 45k. "
         "Trong final không dùng validation/test để chọn epoch; best.pt tương ứng epoch cố định được lưu gần nhất, "
         "không phải validation-best. Thư mục _final tách riêng, giữ nguyên checkpoint selection.")
    para(final_note(runs) + " File final_status_snapshot.json và "
         "final_progress.csv (nếu có) ghi lại tiến độ lúc tổng hợp. Tiến trình train vẫn có thể tiếp tục "
         "sau khi báo cáo đã được tạo; cần tạo lại report khi có kết quả mới.")
    heading("10. Checkpoint, resume, logging và môi trường", True)
    para("Mỗi experiment giữ một last.pt và một best.pt. Selection last lưu model raw, EMA, optimizer, "
         "AMP scaler, RNG Python/NumPy/CPU/CUDA, global_step, history, best weights, cấu hình và split/source hash. "
         "Best chỉ ghi đè khi validation tăng; hòa giữ epoch cũ.")
    para("Checkpoint và log được cập nhật sau epoch train+validation hoàn chỉnh; best ghi trước last khi cần "
         "cập nhật cả hai. Atomic replace, flush/fsync bảo vệ bản đã lưu. Chỉ in một summary mỗi epoch; "
         "không log theo batch. Khi dừng giữa epoch, lần chạy sau khôi phục last và chạy lại epoch dở.")
    para("Resume từ chối khi cấu hình, mã nguồn, phiên bản checkpoint, split hoặc runtime khác. "
         "Dataset sử dụng seed riêng theo sample+epoch và DataLoader seed để replay. "
         "Reproducibility không tự động bảo đảm giữa phần cứng/thư viện khác nhau.")
    environment = runs["cifar10"]["result"]["metadata"]["environment"]
    table(["Thành phần", "Phiên bản / giá trị"], [[key, value] for key, value in environment.items()])
    para("Python duy nhất: C:\\Users\\anhna\\AppData\\Local\\Microsoft\\WindowsApps\\python3.13.exe "
         "(hoặc executable resolved của cùng Store installation). sys.prefix==sys.base_prefix đã kiểm tra. "
         "Không tạo môi trường mới, không tải thư viện hoặc mô hình ngoài khi xây báo cáo.")
    para("Lệnh train/resume: python -u train.py. Lệnh này xử lý selection hai dataset, rồi retrain final/test. "
         "Best test 45k trong báo cáo được đo riêng; không ghi đè test_accuracy=null của result selection "
         "thành kết quả final. Hash SHA256 của weights được kiểm tra khớp file test_evaluation.json.")
    table(["Dataset", "SHA256 của best.pt đã test"], [[NAMES[d], runs[d]["sha256"]] for d in DATASETS])
    heading("11. Code convolution và forward thực tế", True)
    para("Các đoạn sau được trích tự động từ source hiện hành đã khớp source fingerprint của hai run. "
         "Mã đầy đủ ở code/src/models.py, cùng engine.py, datasets.py, recipes.py và các module liên quan. "
         "Không phải pseudo-code viết lại khác implementation.")
    for symbol in ("MixBlock", "ChannelGate", "CompactMixNet", "DropPath", "GlobalPool"):
        code("src/models.py", symbol)
    heading("11.1. Code train và đánh giá", True)
    code("src/engine.py", "train_epoch")
    code("src/engine.py", "evaluate")
    para("Các forward khác được lưu đầy đủ trong models.py: PreActBlock.forward, ResidualNet.forward, "
         "ConvNeXtBlock.forward, CIFARConvNeXt.forward. Adapter các family khác ở timm_models.py. "
         "Không gán kết quả CompactMix-L cho các model đó.")
    heading("12. Chuẩn bị vấn đáp", True)
    questions = [
        ("Vì sao dùng Conv1×1?", "Đổi số kênh và trộn thông tin kênh, không trực tiếp mở rộng receptive field không gian."),
        ("Depthwise khác grouped thế nào?", "Depthwise có groups bằng số kênh đầu vào; grouped chia một số nhóm ít hơn. Block dùng cả hai."),
        ("Vì sao hai kernel 3×3 và 5×5?", "Kết hợp hai thang ngữ cảnh với chi phí depthwise thấp; mức tăng accuracy riêng chưa được ablation xác nhận."),
        ("Residual giúp gì?", "Tạo đường identity cho thông tin và gradient; cộng được vì nhánh residual trả đúng shape C×H×W."),
        ("SE tính attention ở đâu?", "Theo kênh từ global pooling; không phải self-attention theo token."),
        ("Vì sao không đưa softmax vào forward?", "Loss dùng log_softmax ổn định và top-1 chỉ cần argmax logits."),
        ("Vì sao có EMA?", "Theo dõi trung bình trọng số; phải đo validation để chọn, không mặc định luôn tốt hơn raw."),
        ("Vì sao train accuracy không bằng accuracy sạch?", "Mixup/CutMix dùng hai nhãn và metric có trọng số lambda."),
        ("SGD so với Adam?", "SGD dùng momentum gradient; Adam dùng moment bậc một/hai thích nghi. Chưa có full-data Adam để so sánh bằng số."),
        ("Validation và test khác gì?", "Validation dùng chọn mô hình; test dùng đánh giá checkpoint đã chọn. Tối ưu dựa test làm mất ý nghĩa held-out."),
        ("50k retrain có giữ checkpoint 45k không?", "Có, selection và _final là hai thư mục riêng; final được khởi tạo từ đầu với epoch count freeze."),
        ("Dừng train có mất tiến độ không?", "Giữ last của epoch hoàn chỉnh, có optimizer/EMA/scaler/RNG; epoch dở chạy lại."),
        ("Đã hoàn thành toàn bộ đề chưa?", "Chưa: còn so sánh optimizer/LR. "
         + final_note(runs) + " Hai selection/test 45k đã hoàn thành.")]
    for question, answer in questions:
        para(question + " — " + answer)
    heading("13. Tài liệu và chỉ dẫn kiểm chứng", True)
    para("Nguồn số liệu: evidence/cifar10/ và evidence/cifar100/. File test_evaluation.json là chứng cứ "
         "test trực tiếp; metrics.csv và training.log là lịch sử epoch; split.json là index dữ liệu; "
         "layers.csv và model_repr.txt mô tả model; manifest.json liệt kê SHA256 các file trong gói.")
    para("Nguồn đề bài: evidence/Final exam.docx. Source triển khai: code/train.py và code/src/*.py. "
         "Các paper dưới đây giải thích thành phần, không phải baseline được đo trong báo cáo.")
    for reference in [
        "CIFAR datasets: https://www.cs.toronto.edu/~kriz/cifar.html",
        "PyTorch Conv2d: https://docs.pytorch.org/docs/stable/generated/torch.nn.Conv2d.html",
        "PyTorch randomness: https://docs.pytorch.org/docs/stable/notes/randomness.html",
        "Residual learning: He et al., 2015, https://arxiv.org/abs/1512.03385",
        "Squeeze-and-Excitation: Hu et al., 2017, https://arxiv.org/abs/1709.01507",
        "Mixup: Zhang et al., 2017, https://arxiv.org/abs/1710.09412",
        "CutMix: Yun et al., 2019, https://arxiv.org/abs/1905.04899",
        "AutoAugment: Cubuk et al., 2018, https://arxiv.org/abs/1805.09501"]:
        para(reference)
    return content


def dataset_section(content, item, dataset):
    result, test, best = item["result"], item["test"]["test_metrics"], item["best_row"]
    last = item["history"][-1]
    content.append(("table", ["Chỉ số", "Giá trị"], [
        ["Epoch đã hoàn thành / best", f'600 / {result["best_epoch"]} ({result["best_kind"]})'],
        ["Best epoch: raw val / EMA val", f'{best["val_accuracy"]:.2f}% / {best["ema_val_accuracy"]:.2f}%'],
        ["Epoch 600: raw val / EMA val", f'{last["val_accuracy"]:.2f}% / {last["ema_val_accuracy"]:.2f}%'],
        ["Epoch 600: train mixed accuracy", f'{last["train_mixed_accuracy"]:.4f}%'],
        ["Test top-1 / top-5", f'{test["accuracy"]:.2f}% / {test["top5_accuracy"]:.2f}%'],
        ["Test đúng / sai", f'{item["test"]["correct"]} / {item["test"]["incorrect"]}'],
        ["Test loss / macro-F1", f'{test["loss"]:.6f} / {test["macro_f1"]:.6f}'],
        ["Thời gian cộng dồn train + val", clock_duration(result["training_time"])],
        ["Thời gian / epoch trung bình", f'{result["training_time"]/600:.2f} giây'],
        ["Tổng AMP skipped updates", sum(r["amp_skips"] for r in item["history"])],
        ["Test precision / batch / augmentation", "float32 / 64 / không augmentation, không TTA"]]))
    content.append(("image", f"figures/{dataset}_training_curves.png",
                    "Đường cong từ đủ 600 epoch. Đường đứt đánh dấu best epoch. Train mixed accuracy và train loss dùng nhãn trộn."))
    content.append(("paragraph", "Thời gian cộng dồn là số engine lưu cho xử lý các epoch, không tính khoảng nghỉ giữa "
                    "các lần chạy. Hai test score được xác thực bằng confusion matrix và SHA256 checkpoint; "
                    "không suy diễn từ validation hoặc từ dữ liệu synthetic."))
    content.append(("image", f"figures/{dataset}_confusion_matrix.png",
                    "Confusion matrix: hàng=nhãn thật, cột=nhãn dự đoán, màu=phần trăm theo hàng. "
                    "CIFAR-10 hiển thị số lượng; CIFAR-100 có bảng 100×100 đầy đủ trong CSV."))
    content.append(("image", f"figures/{dataset}_class_recall.png",
                    "Recall theo lớp từ cùng phép đánh giá test. Precision, recall, F1 và support đầy đủ trong per_class_metrics.csv."))
    weakest = sorted(item["per_class"], key=lambda r: r["recall_percent"])[:5]
    content.append(("table", ["5 lớp recall thấp nhất", "Recall (%)", "Đúng / support"], [
        [r["class_name"], f'{r["recall_percent"]:.2f}', f'{r["correct"]}/{r["support"]}'] for r in weakest]))


def copy_source():
    for relative in ["train.py", "requirements.txt", "requirements-lock.txt"] + [
            str(path.relative_to(ROOT)) for path in sorted((ROOT / "src").glob("*.py"))] + [
            str(path.relative_to(ROOT)) for path in sorted((ROOT / "configs").glob("*.yaml"))]:
        destination = OUT / "code" / relative
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(ROOT / relative, destination)
    shutil.copyfile(ROOT / "Final exam.docx", OUT / "evidence/Final exam.docx")
    with zipfile.ZipFile(ROOT / "Final exam.docx") as archive:
        document = ET.fromstring(archive.read("word/document.xml"))
        namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        paragraphs = ["".join(p.itertext()) for p in document.findall(".//w:p", namespace)]
    (OUT / "evidence/exam_requirements.txt").write_text("\n".join(paragraphs), encoding="utf-8")
    evaluation_script = '''"""Re-evaluate an archived best using the user's selected Python and local CIFAR data."""
import argparse
import json
import os
from pathlib import Path
import sys
HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE / "code"))
from src.python_policy import ensure_selected_python
ensure_selected_python()
os.environ.setdefault("CUBLAS_WORKSPACE_CONFIG", ":4096:8")
import torch
from torchvision import datasets
from src import engine
from src.datasets import make_loader
from src.models import build_model

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--dataset", choices=["cifar10", "cifar100"], required=True)
    parser.add_argument("--data-root", type=Path, default=HERE.parent / "datasets")
    parser.add_argument("--device", choices=["cuda", "cpu"], default="cuda")
    args = parser.parse_args()
    checkpoint = engine.load_checkpoint(HERE / "checkpoints" / args.dataset / "best.pt")
    config = dict(checkpoint["config"])
    config.update(data_root=str(args.data_root.resolve()), device=args.device, batch_size=64, workers=0)
    engine.seed_everything(config["seed"], config["deterministic"])
    torch.set_num_threads(2)
    classes = 10 if args.dataset == "cifar10" else 100
    model = build_model(config["model"], classes)
    model.load_state_dict(checkpoint["model"], strict=True)
    model = model.to(args.device)
    constructor = datasets.CIFAR10 if classes == 10 else datasets.CIFAR100
    source = constructor(root=config["data_root"], train=False, download=False)
    loader = make_loader(source, list(range(len(source))), config, 0, training=False)
    metrics = engine.evaluate(model, loader, args.device, extended=True)
    print(json.dumps({"dataset": args.dataset, "epoch": checkpoint["epoch"],
                      "samples": len(source), "metrics": metrics}, indent=2))

if __name__ == "__main__":
    main()
'''
    (OUT / "evaluate_best.py").write_text(evaluation_script, encoding="utf-8")
    save_json(OUT / "code/source_snapshot.json", {
        "captured_at": STAMP.isoformat(), "training_source_sha256": engine.source_fingerprint(),
        "repository_commit": subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip()})


def render_markdown(content):
    fence = chr(96) * 3
    lines = ["# Báo cáo bài làm: CompactMix-L trên CIFAR-10 và CIFAR-100",
             "", f"Thời điểm tổng hợp: {STAMP:%d/%m/%Y %H:%M:%S} (UTC+7).",
             "", "Bản PDF để nộp; bản DOCX để điền tên/lớp và chỉnh sửa.", ""]
    for item in content:
        if item[0] == "heading":
            lines.extend(["## " + item[1], ""])
        elif item[0] == "paragraph":
            lines.extend([item[1], ""])
        elif item[0] == "table":
            lines.append("| " + " | ".join(map(str, item[1])) + " |")
            lines.append("| " + " | ".join("---" for _ in item[1]) + " |")
            lines.extend("| " + " | ".join(str(cell).replace("|", "/") for cell in row) + " |" for row in item[2])
            lines.append("")
        elif item[0] == "image":
            lines.extend([f"![{item[2]}]({item[1]})", "", item[2], ""])
        elif item[0] == "code":
            lines.extend(["### " + item[1], "", fence + "python", item[2], fence, ""])
    (OUT / "Bao_cao_CIFAR10_CIFAR100.md").write_text("\n".join(lines), encoding="utf-8")


def render_pdf(content, runs):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak, Preformatted
    for name, filename in [("VN", "arial.ttf"), ("VN-Bold", "arialbd.ttf"), ("Code", "consola.ttf")]:
        pdfmetrics.registerFont(TTFont(name, str(Path("C:/Windows/Fonts") / filename)))
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("VNBody", fontName="VN", fontSize=10, leading=15, spaceAfter=9))
    styles.add(ParagraphStyle("VNTitle", fontName="VN-Bold", fontSize=25, leading=33, alignment=TA_CENTER, spaceAfter=18))
    styles.add(ParagraphStyle("VNHeading", fontName="VN-Bold", fontSize=15, leading=20, textColor=colors.HexColor(BLUE), spaceAfter=13))
    styles.add(ParagraphStyle("VNCaption", fontName="VN", fontSize=8.5, leading=12, textColor=colors.HexColor("#536575"), spaceAfter=12))
    styles.add(ParagraphStyle("VNCell", fontName="VN", fontSize=8.5, leading=12))
    body_width = A4[0] - 3.6 * cm
    doc = SimpleDocTemplate(str(OUT / "Bao_cao_CIFAR10_CIFAR100.pdf"), pagesize=A4,
                            rightMargin=1.8*cm, leftMargin=1.8*cm, topMargin=1.7*cm,
                            bottomMargin=1.7*cm, title="CompactMix-L: CIFAR-10 và CIFAR-100",
                            author="Báo cáo bài làm")
    story = [Spacer(1, 2.5*cm), Paragraph("BÁO CÁO BÀI LÀM", styles["VNTitle"]),
             Paragraph("CompactMix-L<br/>CIFAR-10 &amp; CIFAR-100", styles["VNTitle"]),
             Spacer(1, .5*cm),
             Paragraph("Thiết kế kiến trúc • Huấn luyện • Đánh giá • Mã nguồn", styles["VNBody"]),
             Paragraph(f"Thời điểm số liệu: {STAMP:%d/%m/%Y %H:%M:%S} (UTC+7)", styles["VNBody"]),
             Paragraph("Sinh viên / nhóm: ........................................................", styles["VNBody"]),
             Paragraph("Lớp / mã học phần: ........................................................", styles["VNBody"]),
             Paragraph("Giảng viên: ....................................................................", styles["VNBody"]),
             Spacer(1, 1*cm), Paragraph("Test best train 45k: CIFAR-10 97,18% · CIFAR-100 82,98%.<br/>"
                 + html.escape(final_note(runs)) + " So sánh SGD/Adam và nhiều LR còn thiếu.", styles["VNBody"]), PageBreak()]
    for item in content:
        if item[0] == "heading":
            if item[2]: story.append(PageBreak())
            story.append(Paragraph(html.escape(item[1]), styles["VNHeading"]))
        elif item[0] == "paragraph":
            story.append(Paragraph(html.escape(item[1]), styles["VNBody"]))
        elif item[0] == "table":
            rows = [[Paragraph(html.escape(str(cell)), styles["VNCell"]) for cell in row] for row in [item[1]] + item[2]]
            count = len(item[1])
            widths = [body_width/count] * count
            if count == 2: widths = [body_width*.37, body_width*.63]
            table = Table(rows, colWidths=widths, repeatRows=1, hAlign="LEFT", rowSplitRange=(3, -1))
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCECF4")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#B7C9D3")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F8FA")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
            story.extend([table, Spacer(1, .3*cm)])
        elif item[0] == "image":
            picture = Image(str(OUT / item[1]))
            scale = min(body_width/picture.imageWidth, 21*cm/picture.imageHeight)
            picture.drawWidth = picture.imageWidth*scale
            picture.drawHeight = picture.imageHeight*scale
            story.extend([picture, Paragraph(html.escape(item[2]), styles["VNCaption"])])
        elif item[0] == "code":
            story.append(Paragraph(html.escape(item[1]), styles["VNCaption"]))
            wrapped = "\n".join("\n".join(textwrap.wrap(line, width=93, replace_whitespace=False,
                                  drop_whitespace=False, subsequent_indent="    ")) if len(line) > 93 else line
                                  for line in item[2].splitlines())
            code_style = ParagraphStyle("CodeBlock", fontName="Code", fontSize=7.4, leading=10,
                                        backColor=colors.HexColor("#F1F4F6"), spaceAfter=14)
            story.append(Preformatted(wrapped, code_style))
    def footer(canvas, doc):
        canvas.setFont("VN", 8)
        canvas.setFillColor(colors.HexColor("#536575"))
        canvas.drawString(1.8*cm, .85*cm, "CompactMix-L · Kết quả checkpoint selection 45k · Báo cáo theo bằng chứng")
        canvas.drawRightString(A4[0]-1.8*cm, .85*cm, str(doc.page))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def render_docx(content, runs):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(1.8)
    section.top_margin = section.bottom_margin = Cm(1.7)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"; normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(7)
    document.add_heading("BÁO CÁO BÀI LÀM", 0)
    document.add_heading("CompactMix-L trên CIFAR-10 và CIFAR-100", 1)
    document.add_paragraph(f"Thời điểm số liệu: {STAMP:%d/%m/%Y %H:%M:%S} (UTC+7)")
    for label in ("Sinh viên / nhóm", "Lớp / mã học phần", "Giảng viên"):
        document.add_paragraph(label + ": ........................................................")
    document.add_paragraph("Test best train 45k: CIFAR-10 97,18%; CIFAR-100 82,98%. "
                           + final_note(runs) + " So sánh optimizer/LR còn thiếu.")
    document.add_page_break()
    for item in content:
        if item[0] == "heading":
            if item[2]: document.add_page_break()
            document.add_heading(item[1], 1)
        elif item[0] == "paragraph":
            document.add_paragraph(item[1])
        elif item[0] == "table":
            table = document.add_table(rows=1, cols=len(item[1]))
            table.style = "Light Shading Accent 1"
            for cell, text in zip(table.rows[0].cells, item[1]):
                cell.text = str(text)
            repeat = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in item[2]:
                for cell, text in zip(table.add_row().cells, row):
                    cell.text = str(text)
            document.add_paragraph()
        elif item[0] == "image":
            from PIL import Image
            with Image.open(OUT / item[1]) as picture:
                ratio = picture.height/picture.width
            width = min(17.4, 20.5/ratio)
            document.add_picture(str(OUT / item[1]), width=Cm(width))
            document.add_paragraph(item[2], style="Caption")
        elif item[0] == "code":
            document.add_heading(item[1], 2)
            for line in item[2].splitlines():
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(0)
                run = paragraph.add_run(line or " ")
                run.font.name = "Consolas"; run.font.size = Pt(8)
                run.font.color.rgb = RGBColor.from_string("203849")
    footer = section.footer.paragraphs[0]
    footer.text = "CompactMix-L · Báo cáo bài làm · "
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.save(OUT / "Bao_cao_CIFAR10_CIFAR100.docx")


def supplementary(runs):
    # Keep byte hashes stable across Git checkouts; generated images are reviewed visually.
    (OUT / ".gitattributes").write_text(
        "* -text\n*.pdf binary\n*.docx binary\n*.png binary\n*.svg binary\n*.pt binary\n",
        encoding="utf-8")
    rows = []
    for dataset, item in runs.items():
        rows.append({"dataset": NAMES[dataset], "phase": "selection_45k",
                     "train_images": 45000, "validation_images": 5000, "test_images": 10000,
                     "epochs": 600, "best_epoch": item["result"]["best_epoch"],
                     "best_kind": item["result"]["best_kind"],
                     "validation_accuracy": item["result"]["best_val_accuracy"],
                     "test_accuracy": item["test"]["test_metrics"]["accuracy"],
                     "test_top5_accuracy": item["test"]["test_metrics"]["top5_accuracy"],
                     "test_macro_f1": item["test"]["test_metrics"]["macro_f1"],
                     "parameters": item["profile"]["parameters"], "flops": item["profile"]["flops"],
                     "optimizer": "sgd", "initial_lr": .1, "checkpoint_sha256": item["sha256"]})
    write_csv(OUT / "results_summary.csv", rows)
    save_json(OUT / "results_summary.json", {"captured_at": STAMP.isoformat(), "selection_results": rows,
              "final_status": {d: runs[d]["final"] for d in DATASETS},
              "missing_exam_comparisons": ["SGD additional initial LRs", "Adam", "multiple seeds", "ablations"]})
    readme = f"""# Gói báo cáo CIFAR-10 và CIFAR-100

Mở **Bao_cao_CIFAR10_CIFAR100.pdf** để đọc/nộp. Bản **DOCX** dùng để điền tên, lớp, giảng viên và chỉnh sửa.
Bản **Markdown** cho phép xem trực tiếp trong IDE.
Snapshot số liệu: {STAMP:%d/%m/%Y %H:%M:%S} (UTC+7).

| Dataset | Best validation | Test best train 45k | Best epoch |
|---|---:|---:|---:|
| CIFAR-10 | 97,56% | 97,18% | 567 (EMA) |
| CIFAR-100 | 83,36% | 82,98% | 399 (EMA) |

## Thành phần

- figures/: sơ đồ tổng thể, MixBlock, SE, giao thức dữ liệu; learning curves, confusion matrix, recall theo lớp. Có PNG và SVG.
- evidence/cifar10/ và evidence/cifar100/: config, metadata, 600 epoch CSV/log, kết quả selection, split đầy đủ, test JSON, layer shapes/parameters, confusion matrix và per-class CSV.
- evidence/Final exam.docx: đề bài gốc.
- code/: bản sao source train, Conv/forward, data, optimizer, loss, EMA, resume, final evaluation, search và config.
- checkpoints/cifar10/best.pt và checkpoints/cifar100/best.pt: hai trọng số best 45k đã test. Không sao chép last.pt; bản resume vẫn nằm ở checkpoints/ của repository.
- evaluate_best.py: chạy lại đánh giá archived best trên dữ liệu CIFAR local; không tải dataset và không train.
- results_summary.csv/json: bảng chỉ số gọn.
- manifest.json: SHA256 toàn bộ file trong gói và các kiểm tra tính đúng.

## Đọc kết quả đúng phạm vi

Hai số test trên là kết quả best train trên 45k. {final_note(runs)}
Source selection result.json giữ test_accuracy=null theo giao thức gốc; test đo riêng nằm ở test_evaluation.json.
Đề còn yêu cầu so sánh SGD/Adam và nhiều initial LR: hiện mới có full-data SGD, LR=0,1. Báo cáo nêu rõ phần chưa hoàn thành.
Không coi hỗ trợ --matrix hay Optuna là kết quả đã train; không tuyên bố ablation/multi-seed đã hoàn tất.

## Lệnh dùng môi trường Python duy nhất của người dùng

Từ thư mục gốc D:\\DeepLearning\\Finnal:

    python scripts/build_submission_report.py
    python report/evaluate_best.py --dataset cifar10 --data-root datasets
    python report/evaluate_best.py --dataset cifar100 --data-root datasets

Alias python chỉ dùng khi đã xác minh là cùng Microsoft Store Python 3.13.
Executable chính thức:
C:\\Users\\anhna\\AppData\\Local\\Microsoft\\WindowsApps\\python3.13.exe

Không cần tạo hoặc cài thêm môi trường. Việc build report không train, không thay đổi checkpoint nguồn.
Weights *.pt có trong thư mục cục bộ nhưng bị Git ignore theo chính sách repository; khi gửi folder phải gửi kèm hai file này.
"""
    (OUT / "README.md").write_text(readme, encoding="utf-8")


def verify_and_manifest(runs):
    import pymupdf
    from docx import Document
    from PIL import Image
    pdf_path = OUT / "Bao_cao_CIFAR10_CIFAR100.pdf"
    with pymupdf.open(pdf_path) as document:
        pdf_text = "\n".join(page.get_text() for page in document)
        pages = len(document)
        assert pages >= 15
        for token in ("97,18%", "82,98%", "97,56%", "83,36%", "CHƯA HOÀN THÀNH", "MixBlock"):
            assert token in pdf_text, f"PDF missing {token}"
        preview = OUT / "figures" / "preview_pdf"
        preview.mkdir(exist_ok=True)
        for index in range(pages):
            document[index].get_pixmap(matrix=pymupdf.Matrix(1, 1)).save(preview / f"page_{index+1:02d}.png")
    word = Document(OUT / "Bao_cao_CIFAR10_CIFAR100.docx")
    text = "\n".join(p.text for p in word.paragraphs) + "\n".join(
        cell.text for table in word.tables for row in table.rows for cell in row.cells)
    for token in ("97,18%", "82,98%", "97,56%", "83,36%", "MixBlock"):
        assert token in text
    for picture in (OUT / "figures").rglob("*.png"):
        with Image.open(picture) as file:
            file.verify()
    assert engine.source_fingerprint() == runs["cifar10"]["result"]["metadata"]["source_sha256"]
    files = {}
    for path in sorted(OUT.rglob("*")):
        if path.is_file() and path.name != "manifest.json":
            files[path.relative_to(OUT).as_posix()] = {"sha256": sha(path.read_bytes()), "bytes": path.stat().st_size}
    for dataset in DATASETS:
        assert files[f"checkpoints/{dataset}/best.pt"]["sha256"] == runs[dataset]["sha256"]
        assert np.trace(runs[dataset]["matrix"]) == runs[dataset]["test"]["correct"]
    save_json(OUT / "manifest.json", {"captured_at": STAMP.isoformat(),
              "source_sha256": engine.source_fingerprint(), "pdf_pages": pages, "file_count": len(files),
              "verified": ["best weights match test SHA256", "10k confusion counts match top1",
                           "600 contiguous epoch histories", "45k/5k disjoint full split",
                           "parameter count and CPU forward profile match",
                           "PDF and DOCX contain measured metrics", "PNG decoding", "source unchanged"],
              "files": files})
    return pages, len(files)


def main():
    OUT.mkdir(exist_ok=True)
    runs = collect()
    print("Evidence verified: both 600-epoch runs, best hashes, 10k test counts and model profiles.", flush=True)
    copy_source()
    architecture_diagrams()
    result_figures(runs)
    content = make_content(runs)
    render_markdown(content)
    render_pdf(content, runs)
    render_docx(content, runs)
    supplementary(runs)
    pages, files = verify_and_manifest(runs)
    print(json.dumps({"report_folder": str(OUT), "pdf_pages": pages, "artifact_files": files,
                      "test_accuracy": {d: runs[d]["test"]["test_metrics"]["accuracy"] for d in DATASETS},
                      "final_snapshot": {d: {k:v for k,v in runs[d]["final"].items() if k != "freeze"} for d in DATASETS}},
                     ensure_ascii=False, indent=2), flush=True)


if __name__ == "__main__":
    main()
